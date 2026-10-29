import openpyxl,re,fnmatch,sys
from openpyxl.utils import range_boundaries, get_column_letter
P=sys.argv[1]
class NAt: pass
NA=NAt()
REF=r"(?:'([^']+)'!|([A-Za-z0-9_]+)!)?\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?"
def run(sel):
    wb=openpyxl.load_workbook(P); bp=wb['Build your programme']
    for r in range(1,60):
        if bp.cell(r,5).value in('Yes','No'): bp.cell(r,5).value=sel.get(bp.cell(r,1).value,'No')
    cache={}; errs=[]
    def val(sh,a):
        k=(sh,a)
        if k in cache: return cache[k]
        v=wb[sh][a].value
        if isinstance(v,str) and v.startswith('='): v=ev(sh,v[1:])
        cache[k]=v; return v
    def rng(sh,a,b):
        c1,r1,c2,r2=range_boundaries(f'{a}:{b}')
        return [val(sh,f'{get_column_letter(c)}{r}') for r in range(r1,r2+1) for c in range(c1,c2+1)]
    def topsplit(s):
        out=[];d=0;q=False;cur=''
        for ch in s:
            if ch=='"': q=not q
            if not q:
                if ch=='(': d+=1
                if ch==')': d-=1
                if ch=='&' and d==0: out.append(cur);cur='';continue
            cur+=ch
        return out+[cur]
    G=dict(TIME=lambda h,m,s:(h*3600+m*60+s)/86400, IF=lambda c,a,b=False:a if c else b, AND=lambda *a:all(a),
      COUNTIF=lambda r,c:sum(1 for x in r if x==c), MATCH=lambda v,r,t:(r.index(v)+1) if v in r else NA,
      INDEX=lambda r,i:NA if i is NA else r[i-1], IFERROR=lambda v,a:a if v is NA else v, ISNA=lambda v:v is NA,
      CAT=lambda *a:''.join(str(int(x)) if isinstance(x,(int,float)) and float(x).is_integer() else str(x) for x in a))
    def ev(sh,f):
        parts=topsplit(f)
        if len(parts)>1: f='CAT('+','.join(parts)+')'
        def sub(m):
            s=m.group(1) or m.group(2) or sh
            if m.group(5): return f'rng({s!r},"{m.group(3)}{m.group(4)}","{m.group(5)}{m.group(6)}")'
            return f'val({s!r},"{m.group(3)}{m.group(4)}")'
        py=''.join(p if p.startswith('"') else re.sub(r'(?<![A-Za-z])'+REF+r'(?![\(A-Za-z])',sub,p).replace('<>','!=').replace('=','==').replace('!==','!=').replace('<==','<=').replace('>==','>=') for p in re.split(r'("[^"]*")',f))
        try: return eval(py,dict(G,val=val,rng=rng))
        except Exception as e: errs.append((sh,f,repr(e))); return None
    n=0
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value,str) and c.value.startswith('='): val(ws.title,c.coordinate); n+=1
    return wb,val,errs,n
for sel in [{}, {'Pre-course':'Yes','Add-on 1':'Yes','Add-on 6':'Yes'}, {k:'Yes' for k in ['Pre-course']+[f'Add-on {i}' for i in range(1,14)]}]:
    wb,val,errs,n=run(sel); bp=wb['Build your programme']
    print('\nSEL',sel,'formulas',n,'errors',len(errs),errs[:3])
    for r in range(1,70):
        b=bp.cell(r,2).value
        if isinstance(b,str) and b.startswith('='):
            v=val(bp.title,f'B{r}')
            if v: print(' ',val(bp.title,f'A{r}') if str(bp.cell(r,1).value).startswith('=') else bp.cell(r,1).value,'|',v,'|',val(bp.title,f'C{r}') if str(bp.cell(r,3).value).startswith('=') else '')
hm=lambda v:f'{round(v*1440)//60:02d}:{round(v*1440)%60:02d}'
for sh in []:
    ws=wb[sh]; 
    for r in range(5,ws.max_row+1):
        a=ws.cell(r,1).value
        if a and (str(a)[:1]=='S' and str(a)[1:].isdigit() or a in('Break','Lunch')): print(sh,a,hm(val(sh,f'B{r}')),hm(val(sh,f'C{r}')),(ws.cell(r,4).value or '')[:40])
        elif a and len(str(a))>12: print(sh,'##',str(a)[:60])
